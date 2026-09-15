import csv
from io import BytesIO, StringIO
import json

from django.core.exceptions import ValidationError as DjangoValidationError
from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from rest_framework.exceptions import ValidationError

from fashionerp.audit.services import audit_snapshot, record_audit_event
from fashionerp.authorization.services import has_any_scope_permission, has_permission

from .custom_fields import active_custom_fields, normalize_custom_values, save_custom_values, visible_custom_values
from .data_resources import get_data_resource
from .security import effective_field_edit_permission


MAX_IMPORT_ROWS = 10_000
MAX_IMPORT_BYTES = 10 * 1024 * 1024
IMPORT_MODES = {"create", "update", "upsert", "skip_duplicates"}


def parse_mapping(value) -> dict[str, str]:
    if isinstance(value, dict):
        mapping = value
    else:
        try:
            mapping = json.loads(value or "{}")
        except (TypeError, ValueError) as exc:
            raise ValidationError({"mapping": "Mapping must be a JSON object."}) from exc
    if not isinstance(mapping, dict) or not mapping:
        raise ValidationError({"mapping": "At least one source column must be mapped."})
    normalized = {str(source).strip(): str(target).strip() for source, target in mapping.items()}
    if any(not source or not target for source, target in normalized.items()):
        raise ValidationError({"mapping": "Source and target field names cannot be empty."})
    targets = list(normalized.values())
    if len(targets) != len(set(targets)):
        raise ValidationError({"mapping": "Two source columns cannot map to the same target field."})
    return normalized


def _validate_headers(headers):
    nonempty = [header for header in headers if header]
    if len(nonempty) != len(set(nonempty)):
        raise ValidationError({"file": "Column headers must be unique."})


def parse_uploaded_table(upload) -> list[dict]:
    if upload.size and upload.size > MAX_IMPORT_BYTES:
        raise ValidationError({"file": "Import file exceeds the 10 MB foundation limit."})
    payload = upload.read()
    filename = (upload.name or "").lower()
    if filename.endswith(".xlsx"):
        workbook = load_workbook(filename=BytesIO(payload), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        except StopIteration:
            return []
        _validate_headers(headers)
        rows = []
        for values in iterator:
            if all(value in (None, "") for value in values):
                continue
            rows.append({headers[index]: value for index, value in enumerate(values) if index < len(headers) and headers[index]})
            if len(rows) > MAX_IMPORT_ROWS:
                raise ValidationError({"file": f"Import exceeds {MAX_IMPORT_ROWS} rows."})
        return rows

    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationError({"file": "CSV files must be UTF-8 encoded."}) from exc
    reader = csv.DictReader(StringIO(text))
    if not reader.fieldnames:
        return []
    _validate_headers([str(item).strip() if item is not None else "" for item in reader.fieldnames])
    rows = []
    for row in reader:
        if not any(value not in (None, "") for value in row.values()):
            continue
        rows.append(dict(row))
        if len(rows) > MAX_IMPORT_ROWS:
            raise ValidationError({"file": f"Import exceeds {MAX_IMPORT_ROWS} rows."})
    return rows


def _row_payload(row: dict, mapping: dict) -> tuple[dict, dict]:
    native = {}
    custom = {}
    for source, target in mapping.items():
        if source not in row:
            raise ValidationError({"mapping": f"Source column is missing: {source}."})
        value = row[source]
        if target.startswith("x_"):
            custom[target] = value
        else:
            native[target] = value
    return native, custom


def _lookup_existing(adapter, user, raw_native):
    missing = [field for field in adapter.identity_fields if raw_native.get(field) in (None, "")]
    if missing:
        return None, missing
    lookup = {field: raw_native[field] for field in adapter.identity_fields}
    return adapter.queryset(user, adapter.manifest.manage_permission).filter(**lookup).first(), []


def _scope_from_validated(validated_data, instance=None):
    company = validated_data.get("company", getattr(instance, "company", None))
    establishment = validated_data.get("establishment", getattr(instance, "establishment", None))
    return company, establishment


def _validate_custom_permissions(user, definitions, custom_values, company, establishment, manifest):
    forbidden = []
    for key in custom_values:
        definition = definitions[key]
        if not has_permission(
            user,
            effective_field_edit_permission(definition, manifest),
            company=company,
            establishment=establishment,
        ):
            forbidden.append(key)
    if forbidden:
        raise DjangoValidationError(
            {"custom": f"Not allowed to edit fields: {', '.join(sorted(forbidden))}."}
        )


def prepare_import(*, request, model_key: str, rows: list[dict], mapping: dict, mode: str) -> list[dict]:
    if mode not in IMPORT_MODES:
        raise ValidationError({"mode": f"Mode must be one of: {', '.join(sorted(IMPORT_MODES))}."})
    if not has_any_scope_permission(request.user, "platform.data.import"):
        raise ValidationError({"permission": "You do not have permission to import data."})
    try:
        adapter = get_data_resource(model_key)
    except LookupError as exc:
        raise ValidationError({"resource": str(exc)}) from exc

    definitions = {item.key: item for item in active_custom_fields(request.user.organization, model_key)}
    allowed = set(adapter.import_fields) | set(definitions)
    invalid_targets = sorted(set(mapping.values()) - allowed)
    if invalid_targets:
        raise ValidationError({"mapping": f"Fields are not importable: {', '.join(invalid_targets)}."})

    plans = []
    for index, row in enumerate(rows, start=2):
        native, custom = _row_payload(row, mapping)
        existing, missing_identity = _lookup_existing(adapter, request.user, native)
        if missing_identity and mode in {"update", "upsert", "skip_duplicates"}:
            plans.append({"row": index, "errors": {"identity": f"Missing identity fields: {', '.join(missing_identity)}."}})
            continue
        if mode == "create" and existing is not None:
            plans.append({"row": index, "errors": {"identity": "A record with this identity already exists."}})
            continue
        if mode == "update" and existing is None:
            plans.append({"row": index, "errors": {"identity": "No existing record matches this identity."}})
            continue
        if mode == "skip_duplicates" and existing is not None:
            plans.append({"row": index, "action": "skip", "native": native, "custom": custom, "instance": existing})
            continue

        action = "update" if existing is not None else "create"
        serializer = adapter.serializer_class(
            instance=existing,
            data=native,
            partial=action == "update",
            context={"request": request},
        )
        if not serializer.is_valid():
            plans.append({"row": index, "errors": serializer.errors})
            continue
        company, establishment = _scope_from_validated(serializer.validated_data, existing)
        if not has_permission(
            request.user,
            adapter.manifest.manage_permission,
            company=company,
            establishment=establishment,
        ):
            plans.append({"row": index, "errors": {"permission": "Record is outside your writable scope."}})
            continue
        try:
            normalized_custom = normalize_custom_values(
                organization=request.user.organization,
                model_key=model_key,
                values=custom,
                partial=action == "update",
            )
            _validate_custom_permissions(
                request.user,
                definitions,
                normalized_custom,
                company,
                establishment,
                adapter.manifest,
            )
        except DjangoValidationError as exc:
            errors = exc.message_dict if hasattr(exc, "message_dict") else {"custom": exc.messages}
            plans.append({"row": index, "errors": errors})
            continue
        plans.append(
            {
                "row": index,
                "action": action,
                "native": native,
                "custom": normalized_custom,
                "instance": existing,
            }
        )
    return plans


def execute_import(*, request, model_key: str, plans: list[dict]) -> dict:
    if any(plan.get("errors") for plan in plans):
        raise ValidationError({"rows": "Import contains validation errors and was not committed."})
    adapter = get_data_resource(model_key)
    created = updated = skipped = 0
    with transaction.atomic():
        for plan in plans:
            if plan["action"] == "skip":
                skipped += 1
                continue
            instance = plan["instance"]
            serializer = adapter.serializer_class(
                instance=instance,
                data=plan["native"],
                partial=plan["action"] == "update",
                context={"request": request},
            )
            serializer.is_valid(raise_exception=True)
            company, establishment = _scope_from_validated(serializer.validated_data, instance)
            if not has_permission(
                request.user,
                adapter.manifest.manage_permission,
                company=company,
                establishment=establishment,
            ):
                raise ValidationError({"permission": "Record scope changed during import."})
            before = audit_snapshot(instance) if instance is not None else None
            if plan["action"] == "create":
                instance = serializer.save(organization=request.user.organization)
                created += 1
            else:
                instance = serializer.save()
                updated += 1
            if plan["custom"]:
                save_custom_values(
                    organization=request.user.organization,
                    model_key=model_key,
                    object_id=instance.pk,
                    values=plan["custom"],
                    actor=request.user,
                    request=request,
                    partial=plan["action"] == "update",
                )
            record_audit_event(
                organization=request.user.organization,
                actor=request.user,
                action=f"platform.data_import.{model_key}.{plan['action']}",
                object_instance=instance,
                before=before,
                after=audit_snapshot(instance),
                request=request,
                metadata={"source_row": plan["row"]},
            )
        record_audit_event(
            organization=request.user.organization,
            actor=request.user,
            action="platform.data_import.commit",
            object_type="platform.data_import",
            object_label=model_key,
            request=request,
            metadata={"created": created, "updated": updated, "skipped": skipped},
        )
    return {"created": created, "updated": updated, "skipped": skipped}


def _json_cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def export_rows(*, request, model_key: str, fields: list[str]) -> list[list]:
    if not has_any_scope_permission(request.user, "platform.data.export"):
        raise ValidationError({"permission": "You do not have permission to export data."})
    adapter = get_data_resource(model_key)
    definitions = {item.key: item for item in active_custom_fields(request.user.organization, model_key)}
    allowed = set(adapter.export_fields) | set(definitions)
    invalid = sorted(set(fields) - allowed)
    if invalid:
        raise ValidationError({"fields": f"Fields are not exportable: {', '.join(invalid)}."})
    rows = []
    queryset = adapter.queryset(request.user, adapter.manifest.view_permission)
    for instance in queryset.iterator():
        custom = visible_custom_values(user=request.user, model_key=model_key, object_id=instance.pk)
        values = []
        for field in fields:
            values.append(_json_cell(custom.get(field, "")) if field.startswith("x_") else _json_cell(getattr(instance, field, "")))
        rows.append(values)
    return rows


def tabular_response(*, headers: list[str], rows: list[list], output_format: str, filename: str):
    if output_format == "csv":
        buffer = StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(rows)
        response = HttpResponse(buffer.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
        return response
    if output_format == "xlsx":
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("Export")
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        workbook.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
        return response
    raise ValidationError({"format": "Format must be csv or xlsx."})
